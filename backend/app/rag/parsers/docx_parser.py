import logging
from pathlib import Path
from typing import Any, List, Union

from app.rag.parsers.base_parser import BaseParser
from app.rag.parsers.unified_document import UnifiedDocument

logger = logging.getLogger("careos.rag.parsers.docx")


class DocxParser(BaseParser):
    """Parser for Microsoft Word (.docx) documents preserving paragraphs, headings, and tables."""

    def parse(
        self,
        file_path: Union[str, Path],
        max_section_length: int = 2500,
        **kwargs: Any,
    ) -> List[UnifiedDocument]:
        """Parse DOCX file and return structured document pages/sections.

        Args:
            file_path: Path to DOCX file.
            max_section_length: Target maximum text length per section before rolling over to a new page index.

        Returns:
            List of UnifiedDocument instances.
        """
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"DOCX file not found at: {file_path}")

        if path.suffix.lower() != ".docx":
            raise ValueError(f"File '{file_path}' is not a valid .docx file.")

        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError(
                f"python-docx is not installed ({e}). Please install python-docx to parse Word documents."
            ) from e

        try:
            doc = Document(str(path))
        except Exception as e:
            raise ValueError(f"Failed to read DOCX file '{file_path}': {str(e)}") from e

        content_blocks: List[str] = []

        # Iterate through paragraphs and tables in natural order
        for block in self._iter_block_items(doc):
            if isinstance(block, str) and block.strip():
                content_blocks.append(block.strip())

        if not content_blocks:
            return []

        # Group blocks into logical pages/sections
        documents: List[UnifiedDocument] = []
        current_page_blocks: List[str] = []
        current_page_len = 0
        current_page_num = 1

        for block in content_blocks:
            block_len = len(block)
            if current_page_blocks and (current_page_len + block_len > max_section_length):
                page_text = "\n\n".join(current_page_blocks).strip()
                if page_text:
                    documents.append(
                        UnifiedDocument(
                            text=page_text,
                            source=path.name,
                            page=current_page_num,
                            source_type="docx",
                            ocr_applied=False,
                            metadata={"extractor": "python-docx", "block_count": len(current_page_blocks)},
                        )
                    )
                    current_page_num += 1
                    current_page_blocks = []
                    current_page_len = 0

            current_page_blocks.append(block)
            current_page_len += block_len

        if current_page_blocks:
            page_text = "\n\n".join(current_page_blocks).strip()
            if page_text:
                documents.append(
                    UnifiedDocument(
                        text=page_text,
                        source=path.name,
                        page=current_page_num,
                        source_type="docx",
                        ocr_applied=False,
                        metadata={"extractor": "python-docx", "block_count": len(current_page_blocks)},
                    )
                )

        return documents

    def _iter_block_items(self, doc: Any):
        """Yield paragraphs and tables in document order."""
        for child in doc.element.body:
            tag = child.tag
            if tag.endswith("p"):  # Paragraph
                for p in doc.paragraphs:
                    if p._element is child:
                        text = p.text.strip()
                        if text:
                            # Format headings clearly
                            if p.style and p.style.name and p.style.name.startswith("Heading"):
                                yield f"## {text}"
                            else:
                                yield text
                        break
            elif tag.endswith("tbl"):  # Table
                for tbl in doc.tables:
                    if tbl._element is child:
                        table_md = self._format_table_as_markdown(tbl)
                        if table_md:
                            yield table_md
                        break

    def _format_table_as_markdown(self, table: Any) -> str:
        """Convert a python-docx table to a formatted Markdown table string."""
        rows_data = []
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            if any(cells):
                rows_data.append(cells)

        if not rows_data:
            return ""

        # Normalize column counts
        max_cols = max(len(r) for r in rows_data)
        normalized_rows = [r + [""] * (max_cols - len(r)) for r in rows_data]

        header = normalized_rows[0]
        separator = ["---"] * max_cols

        md_lines = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(separator) + " |",
        ]

        for row in normalized_rows[1:]:
            md_lines.append("| " + " | ".join(row) + " |")

        return "\n".join(md_lines)

