import base64
import io
import subprocess
from pathlib import Path
from unittest.mock import MagicMock, patch
from uuid import uuid4

import pymupdf as fitz
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image

from app.main import app
from app.rag.loader import (
    extract_document_bytes,
    get_ocr_status_info,
    load_directory,
    load_document,
    load_pdf,
)
from app.rag.ocr import HandwritingOCREngine, OCREngine, OCRRouter, get_ocr_router
from app.rag.parsers import (
    DocxParser,
    ImageParser,
    PDFParser,
    UnifiedDocument,
    detect_document_format,
    get_parser_for_file,
    load_unified_documents,
)


def _register_clinician_with_patient(client: TestClient) -> tuple[dict[str, str], str]:
    owner_email = f"owner-ocr-{uuid4().hex[:8]}@example.org"
    owner = client.post(
        "/api/v1/auth/register",
        json={
            "email": owner_email,
            "password": "strong-password",
            "full_name": "OCR Workspace Owner",
            "organization_name": "CareOS OCR Clinic",
        },
    )
    assert owner.status_code == 201
    owner_headers = {"Authorization": f"Bearer {owner.json()['access_token']}"}

    clinician_email = f"dr-ocr-{uuid4().hex[:8]}@example.org"
    invite = client.post(
        "/api/v1/team/invites",
        headers=owner_headers,
        json={"email": clinician_email, "role": "physician"},
    )
    assert invite.status_code == 201
    clinician = client.post(
        "/api/v1/auth/invitations/accept",
        json={
            "token": invite.json()["development_token"],
            "full_name": "Dr. OCR Tester",
            "password": "strong-password",
        },
    )
    assert clinician.status_code == 200
    headers = {"Authorization": f"Bearer {clinician.json()['access_token']}"}

    patient = client.post(
        "/api/v1/patients",
        headers=headers,
        json={
            "medical_record_number": f"MRN-OCR-{uuid4().hex[:6].upper()}",
            "given_name": "Salma",
            "family_name": "Mansour",
            "date_of_birth": "1985-04-12",
        },
    )
    assert patient.status_code == 201
    return headers, patient.json()["id"]


def _make_digital_pdf_bytes(pages: list[str]) -> bytes:
    doc = fitz.open()
    try:
        for text in pages:
            page = doc.new_page()
            if text:
                page.insert_text((50, 72), text)
        return doc.tobytes()
    finally:
        doc.close()


def _make_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Discharge Summary", level=1)
    doc.add_paragraph("Patient presented with acute bronchitis and mild fever.")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Medication"
    table.rows[0].cells[1].text = "Dose"
    table.rows[0].cells[2].text = "Frequency"
    table.rows[1].cells[0].text = "Amoxicillin"
    table.rows[1].cells[1].text = "500mg"
    table.rows[1].cells[2].text = "TID"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_png_bytes(width: int = 160, height: int = 80) -> bytes:
    img = Image.new("RGB", (width, height), color=(245, 245, 245))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_unified_document_and_format_detection() -> None:
    doc = UnifiedDocument(
        text="Clinical lab result: HbA1c 6.8%",
        source="lab.pdf",
        page=2,
        source_type="pdf",
        ocr_applied=False,
    )
    as_dict = doc.to_dict()
    assert as_dict["source"] == "lab.pdf"
    assert as_dict["page"] == 2
    assert as_dict["metadata"] == {}

    restored = UnifiedDocument.from_dict(as_dict)
    assert restored.text == doc.text
    assert restored.page == 2

    assert detect_document_format("report.pdf") == "pdf"
    assert detect_document_format("notes.DOCX") == "docx"
    assert detect_document_format("scan.png") == "image"
    assert detect_document_format("photo.JPG") == "image"
    assert detect_document_format("rx.jpeg") == "image"
    assert detect_document_format("slip.webp") == "image"
    assert isinstance(get_parser_for_file("a.pdf"), PDFParser)
    assert isinstance(get_parser_for_file("a.docx"), DocxParser)
    assert isinstance(get_parser_for_file("a.png"), ImageParser)

    with pytest.raises(ValueError, match="Unsupported document format"):
        detect_document_format("archive.zip")


def test_real_digital_pdf_and_scanned_pdf_fallback(tmp_path: Path) -> None:
    digital_pdf = tmp_path / "digital_guideline.pdf"
    digital_pdf.write_bytes(
        _make_digital_pdf_bytes(
            [
                "Page 1: Hypertension target BP is < 130/80 mmHg.",
                "",
                "Page 3: Lisinopril 10mg once daily.",
            ]
        )
    )

    parser = PDFParser()
    info = parser.detect_pdf_type(digital_pdf)
    assert info["pdf_type"] == "digital"
    assert info["page_count"] == 3
    assert info["total_text_chars"] > 20

    docs = parser.parse(digital_pdf, enable_ocr_fallback=False)
    assert len(docs) == 2
    assert docs[0].page == 1
    assert "Hypertension target BP" in docs[0].text
    assert docs[0].ocr_applied is False
    assert docs[1].page == 3
    assert "Lisinopril 10mg" in docs[1].text

    legacy_pages = load_pdf(str(digital_pdf))
    assert len(legacy_pages) == 2
    assert legacy_pages[0]["page"] == 1

    # Scanned PDF (blank page with no digital text) routed through mocked OCRRouter
    scanned_pdf = tmp_path / "scanned_rx.pdf"
    scanned_pdf.write_bytes(_make_digital_pdf_bytes([""]))
    scanned_info = parser.detect_pdf_type(scanned_pdf)
    assert scanned_info["pdf_type"] == "scanned"

    mock_router = MagicMock()
    mock_router.extract_text.return_value = {
        "text": "Scanned Prescription: Metformin 850mg BID",
        "confidence": 0.96,
        "ocr_engine": "PaddleOCR",
        "status": "success",
    }
    scanned_docs = parser.parse(scanned_pdf, ocr_router=mock_router)
    assert len(scanned_docs) == 1
    assert scanned_docs[0].ocr_applied is True
    assert scanned_docs[0].ocr_engine == "PaddleOCR"
    assert "Metformin 850mg BID" in scanned_docs[0].text


def test_real_docx_parser_headings_paragraphs_and_tables(tmp_path: Path) -> None:
    docx_file = tmp_path / "clinical_summary.docx"
    docx_file.write_bytes(_make_docx_bytes())

    docs = load_unified_documents(docx_file)
    assert len(docs) == 1
    assert docs[0].source_type == "docx"
    assert "## Discharge Summary" in docs[0].text
    assert "acute bronchitis" in docs[0].text
    assert "| Medication | Dose | Frequency |" in docs[0].text
    assert "| Amoxicillin | 500mg | TID |" in docs[0].text

    loaded_dir = load_directory(tmp_path)
    assert len(loaded_dir) == 1
    assert loaded_dir[0]["source"] == "clinical_summary.docx"


def test_image_parser_and_ocr_engines_lazy_loading(tmp_path: Path) -> None:
    printed = OCREngine(lang="ar", auto_init=False)
    assert printed._engine is None
    handwriting = HandwritingOCREngine(auto_init=False)
    assert handwriting._model is None
    assert handwriting._processor is None

    router = get_ocr_router()
    assert isinstance(router, OCRRouter)
    with pytest.raises(ValueError, match="Unsupported OCR engine type"):
        router.extract_text(Image.new("RGB", (32, 32)), engine_type="invalid_engine")

    img_path = tmp_path / "prescription.png"
    img_path.write_bytes(_make_png_bytes())

    img_parser = ImageParser(apply_preprocessing=True)
    meta = img_parser.validate_image(img_path)
    assert meta["format"] == "PNG"
    assert meta["width"] == 160
    assert meta["height"] == 80

    mock_router = MagicMock()
    mock_router.extract_text.return_value = {
        "text": "Handwritten Rx: Paracetamol 1g q6h",
        "ocr_engine": "Qwen2.5-VL-3B-Handwriting-OCR",
        "confidence": 0.93,
        "status": "success",
    }
    parsed = img_parser.parse(img_path, ocr_engine="handwriting", ocr_router=mock_router)
    assert len(parsed) == 1
    assert parsed[0].ocr_applied is True
    assert parsed[0].ocr_engine == "Qwen2.5-VL-3B-Handwriting-OCR"
    assert "Paracetamol 1g q6h" in parsed[0].text

    # Verify OCREngine.extract_text line aggregation with mocked PaddleOCR
    mock_paddle = MagicMock()
    mock_paddle.predict.return_value = [
        [
            [[[0, 0], [10, 0], [10, 10], [0, 10]], ("CBC Lab Report", 0.98)],
            [[[0, 20], [10, 20], [10, 30], [0, 30]], ("WBC: 7.2 x10^3/uL", 0.94)],
        ]
    ]
    with patch.object(printed, "_get_engine", return_value=mock_paddle):
        res = printed.extract_text(Image.new("RGB", (64, 64)))
        assert res["status"] == "success"
        assert res["ocr_engine"] == "PaddleOCR"
        assert "CBC Lab Report\nWBC: 7.2 x10^3/uL" == res["text"]
        assert round(res["confidence"], 2) == 0.96


def test_printed_ocr_falls_back_to_local_tesseract() -> None:
    printed = OCREngine(lang="ar", auto_init=False)
    tsv = (
        "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
        "5\t1\t1\t1\t1\t1\t0\t0\t20\t10\t95.00\tضغط\n"
        "5\t1\t1\t1\t1\t2\t21\t0\t30\t10\t90.00\tالدم\n"
        "5\t1\t1\t1\t2\t1\t0\t20\t30\t10\t85.00\tStable\n"
    )
    process_result = subprocess.CompletedProcess(
        args=[], returncode=0, stdout=tsv.encode("utf-8"), stderr=b""
    )

    with (
        patch.object(printed, "_get_engine", side_effect=RuntimeError("PaddleOCR is unavailable")),
        patch.object(printed, "_is_tesseract_available", return_value=True),
        patch("app.rag.ocr.ocr_engine.subprocess.run", return_value=process_result) as run,
    ):
        result = printed.extract_text(Image.new("RGB", (64, 64)))

    assert result["text"] == "ضغط الدم\nStable"
    assert result["ocr_engine"] == "Tesseract"
    assert result["language"] == "ar"
    assert round(result["confidence"], 2) == 0.90
    assert run.call_args.args[0] == ["tesseract", "stdin", "stdout", "-l", "ara+eng", "tsv"]
    assert run.call_args.kwargs["input"].startswith(b"\x89PNG")


def test_tesseract_requires_all_configured_language_data() -> None:
    printed = OCREngine(lang="ar", auto_init=False)
    process_result = subprocess.CompletedProcess(
        args=[],
        returncode=0,
        stdout="List of available languages in '/usr/share/tessdata/' (2):\nara\neng\n",
        stderr="",
    )

    with (
        patch("app.rag.ocr.ocr_engine.shutil.which", return_value="/usr/bin/tesseract"),
        patch("app.rag.ocr.ocr_engine.subprocess.run", return_value=process_result),
    ):
        assert printed._is_tesseract_available() is True

    missing_arabic_result = subprocess.CompletedProcess(
        args=[],
        returncode=0,
        stdout="List of available languages (1):\neng\n",
        stderr="",
    )
    with (
        patch("app.rag.ocr.ocr_engine.shutil.which", return_value="/usr/bin/tesseract"),
        patch("app.rag.ocr.ocr_engine.subprocess.run", return_value=missing_arabic_result),
    ):
        assert printed._is_tesseract_available() is False


def test_tesseract_retries_sparse_page_segmentation_when_no_words_are_found() -> None:
    printed = OCREngine(lang="ar", auto_init=False)
    empty_tsv = (
        "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
    )
    recognized_tsv = (
        "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
        "5\t1\t1\t1\t1\t1\t0\t0\t20\t10\t85.00\tضغط\n"
    )

    with (
        patch.object(printed, "_get_engine", side_effect=RuntimeError("PaddleOCR is unavailable")),
        patch.object(printed, "_is_tesseract_available", return_value=True),
        patch(
            "app.rag.ocr.ocr_engine.subprocess.run",
            side_effect=[
                subprocess.CompletedProcess(args=[], returncode=0, stdout=empty_tsv.encode(), stderr=b""),
                subprocess.CompletedProcess(
                    args=[], returncode=0, stdout=recognized_tsv.encode(), stderr=b""
                ),
            ],
        ) as run,
    ):
        result = printed.extract_text(Image.new("RGB", (64, 64)))

    assert result["text"] == "ضغط"
    assert run.call_count == 2
    assert run.call_args_list[1].args[0] == [
        "tesseract",
        "stdin",
        "stdout",
        "-l",
        "ara+eng",
        "--psm",
        "11",
        "tsv",
    ]


def test_inline_rx_directions_are_extracted_separately_from_medication_name() -> None:
    from app.rag.loader import extract_clinical_candidates

    candidates = extract_clinical_candidates(
        "Discharge Summary\nRx: Amlodipine 10mg once daily for 30 days",
        source="synthetic-prescription.txt",
    )

    medication = next(
        medication
        for medication in candidates["medications"]
        if medication["name"].lower() == "amlodipine"
    )
    assert medication["dose"] == "10mg"
    assert medication["frequency"].lower() == "once daily"
    assert medication["duration"].lower() == "30 days"
    assert medication["status"] == "present"


def test_eye_drop_candidates_are_extracted_without_visible_rx_marker() -> None:
    from app.rag.loader import extract_clinical_candidates

    candidates = extract_clinical_candidates(
        "Prescription\n"
        "Efemyo Eye Drop\n"
        "قطرة ٣ مرات يومياً لمدة أسبوع\n"
        "Lubrivic Eye Drop\n"
        "قطرة مرتين يومياً\n"
        "وتكرر عند اللزوم",
        source="synthetic-eye-prescription.jpg",
        ocr_applied=True,
        ocr_engine="Tesseract",
        ocr_confidence=0.56,
    )

    medications = {item["name"]: item for item in candidates["medications"]}
    assert set(medications) == {"Efemyo Eye Drop", "Lubrivic Eye Drop"}
    assert medications["Efemyo Eye Drop"]["dosage"] == "قطرة"
    assert "٣ مرات" in medications["Efemyo Eye Drop"]["frequency"]
    assert "أسبوع" in medications["Efemyo Eye Drop"]["duration"]
    assert medications["Lubrivic Eye Drop"]["duration"] is None
    assert medications["Lubrivic Eye Drop"]["instructions"] == "وتكرر عند اللزوم"
    assert candidates["provenance"]["requires_human_review"] is True
    assert candidates["field_status"]["medications"] == "uncertain"


def test_api_ocr_endpoints_and_patient_document_extraction(monkeypatch) -> None:
    client = TestClient(app)
    headers, patient_id = _register_clinician_with_patient(client)

    # 1. System status and OCR status endpoints
    sys_status = client.get("/api/v1/system/status")
    assert sys_status.status_code == 200
    assert "ocr" in sys_status.json()["providers"]
    assert sys_status.json()["providers"]["ocr"]["configured"] is True

    ocr_status = client.get("/api/v1/ocr/status", headers=headers)
    assert ocr_status.status_code == 200
    status_body = ocr_status.json()
    assert status_body["parsers"]["pdf_pymupdf"] is True
    assert status_body["parsers"]["docx"] is True
    assert status_body["parsers"]["image_pillow"] is True
    assert ".pdf" in status_body["supported_extensions"]
    assert ".docx" in status_body["supported_extensions"]

    # 2. Upload real digital PDF to patient record -> automatic extraction
    pdf_bytes = _make_digital_pdf_bytes(["Cardiology Consult: Ejection fraction 58%, normal sinus rhythm."])
    pdf_upload = client.post(
        f"/api/v1/patients/{patient_id}/documents",
        headers=headers,
        json={
            "filename": "cardiology-consult.pdf",
            "content": base64.b64encode(pdf_bytes).decode(),
            "content_type": "application/pdf",
        },
    )
    assert pdf_upload.status_code == 201
    pdf_data = pdf_upload.json()
    assert pdf_data["ocr_status"] == "completed"
    assert "Ejection fraction 58%" in (pdf_data["extracted_text"] or "")

    # 3. Upload real DOCX to patient record -> automatic extraction
    docx_bytes = _make_docx_bytes()
    docx_upload = client.post(
        f"/api/v1/patients/{patient_id}/documents",
        headers=headers,
        json={
            "filename": "discharge-summary.docx",
            "content": base64.b64encode(docx_bytes).decode(),
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    assert docx_upload.status_code == 201
    docx_data = docx_upload.json()
    assert docx_data["ocr_status"] == "completed"
    assert "Amoxicillin" in (docx_data["extracted_text"] or "")

    # Verify GET /api/v1/patients/{patient_id} returns the extracted_text and ocr_status
    patient_detail = client.get(f"/api/v1/patients/{patient_id}", headers=headers)
    assert patient_detail.status_code == 200
    docs_by_name = {d["filename"]: d for d in patient_detail.json()["documents"]}
    assert docs_by_name["cardiology-consult.pdf"]["ocr_status"] == "completed"
    assert "Ejection fraction 58%" in docs_by_name["cardiology-consult.pdf"]["extracted_text"]
    assert docs_by_name["discharge-summary.docx"]["ocr_status"] == "completed"
    assert "Amoxicillin" in docs_by_name["discharge-summary.docx"]["extracted_text"]

    # 4. Direct /api/v1/ocr/extract endpoint and on-demand /patients/{patient_id}/documents/{doc_id}/ocr
    png_bytes = _make_png_bytes()

    # Upload PNG first without requiring an OCR engine to be installed in the test environment.
    img_upload = client.post(
        f"/api/v1/patients/{patient_id}/documents",
        headers=headers,
        json={
            "filename": "handwritten-rx.png",
            "content": base64.b64encode(png_bytes).decode(),
            "content_type": "image/png",
        },
    )
    assert img_upload.status_code == 201
    img_doc_id = img_upload.json()["id"]

    # Now mock OCRRouter.extract_text to simulate OCR engine response for image extraction
    fake_router = MagicMock()
    fake_router.extract_text.return_value = {
        "text": "Rx: Atorvastatin 20mg once daily at bedtime",
        "ocr_engine": "Qwen2.5-VL-3B-Handwriting-OCR",
        "confidence": 0.95,
        "status": "success",
    }
    monkeypatch.setattr("app.rag.ocr.ocr_router.get_ocr_router", lambda: fake_router)

    direct_extract = client.post(
        "/api/v1/ocr/extract",
        headers=headers,
        json={
            "filename": "handwritten-rx.png",
            "content": base64.b64encode(png_bytes).decode(),
            "content_type": "image/png",
            "ocr_engine": "handwriting",
        },
    )
    assert direct_extract.status_code == 200
    assert direct_extract.json()["ocr_status"] == "completed"
    assert direct_extract.json()["ocr_applied"] is True
    assert direct_extract.json()["ocr_engine"] == "Qwen2.5-VL-3B-Handwriting-OCR"
    assert "Atorvastatin 20mg" in direct_extract.json()["extracted_text"]

    # Trigger on-demand OCR re-processing on the stored patient document
    reprocess = client.post(
        f"/api/v1/patients/{patient_id}/documents/{img_doc_id}/ocr",
        headers=headers,
        json={"ocr_engine": "handwriting"},
    )
    assert reprocess.status_code == 200
    assert reprocess.json()["ocr_status"] == "completed"
    assert "Atorvastatin 20mg" in reprocess.json()["extracted_text"]

    # Confirm updated OCR fields are persisted on the patient document record
    refreshed_patient = client.get(f"/api/v1/patients/{patient_id}", headers=headers)
    assert refreshed_patient.status_code == 200
    refreshed_docs = {d["id"]: d for d in refreshed_patient.json()["documents"]}
    assert refreshed_docs[img_doc_id]["ocr_status"] == "completed"
    assert "Atorvastatin 20mg" in refreshed_docs[img_doc_id]["extracted_text"]

    # Verify unsupported format on /api/v1/ocr/extract returns 422
    bad_extract = client.post(
        "/api/v1/ocr/extract",
        headers=headers,
        json={
            "filename": "unsupported.exe",
            "content": base64.b64encode(b"binary").decode(),
            "content_type": "application/octet-stream",
        },
    )
    assert bad_extract.status_code == 422


def test_ocr_authorization_and_organization_isolation() -> None:
    client = TestClient(app)
    org_a_headers, patient_a1_id = _register_clinician_with_patient(client)
    org_b_headers, _ = _register_clinician_with_patient(client)

    # Create a second patient in Organization A to test cross-patient document mismatch
    patient_a2 = client.post(
        "/api/v1/patients",
        headers=org_a_headers,
        json={
            "medical_record_number": f"MRN-A2-{uuid4().hex[:6].upper()}",
            "given_name": "Kareem",
            "family_name": "Nabil",
            "date_of_birth": "1979-11-20",
        },
    )
    assert patient_a2.status_code == 201
    patient_a2_id = patient_a2.json()["id"]

    pdf_bytes = _make_digital_pdf_bytes(["Org A Private Clinical Document"])
    uploaded = client.post(
        f"/api/v1/patients/{patient_a1_id}/documents",
        headers=org_a_headers,
        json={
            "filename": "org-a-private.pdf",
            "content": base64.b64encode(pdf_bytes).decode(),
            "content_type": "application/pdf",
        },
    )
    assert uploaded.status_code == 201
    doc_a1_id = uploaded.json()["id"]

    # 1. Unauthenticated requests must be rejected with 401
    assert client.get("/api/v1/ocr/status").status_code == 401
    assert (
        client.post(
            "/api/v1/ocr/extract",
            json={
                "filename": "a.pdf",
                "content": base64.b64encode(pdf_bytes).decode(),
                "content_type": "application/pdf",
            },
        ).status_code
        == 401
    )
    assert client.post(f"/api/v1/patients/{patient_a1_id}/documents/{doc_a1_id}/ocr").status_code == 401

    # 2. Non-clinical role (it_admin) must be rejected with 403 on all OCR endpoints
    owner = client.post(
        "/api/v1/auth/register",
        json={
            "email": f"owner-it-{uuid4().hex[:8]}@example.org",
            "password": "strong-password",
            "full_name": "IT Org Owner",
            "organization_name": "IT Org",
        },
    )
    owner_headers = {"Authorization": f"Bearer {owner.json()['access_token']}"}
    it_invite = client.post(
        "/api/v1/team/invites",
        headers=owner_headers,
        json={"email": f"it-admin-{uuid4().hex[:8]}@example.org", "role": "it_admin"},
    )
    it_user = client.post(
        "/api/v1/auth/invitations/accept",
        json={
            "token": it_invite.json()["development_token"],
            "full_name": "IT Admin User",
            "password": "strong-password",
        },
    )
    it_headers = {"Authorization": f"Bearer {it_user.json()['access_token']}"}

    assert client.get("/api/v1/ocr/status", headers=it_headers).status_code == 403
    assert (
        client.post(
            "/api/v1/ocr/extract",
            headers=it_headers,
            json={
                "filename": "a.pdf",
                "content": base64.b64encode(pdf_bytes).decode(),
                "content_type": "application/pdf",
            },
        ).status_code
        == 403
    )
    assert (
        client.post(
            f"/api/v1/patients/{patient_a1_id}/documents/{doc_a1_id}/ocr",
            headers=it_headers,
        ).status_code
        == 403
    )

    # 3. Clinician from Organization B cannot trigger OCR on Organization A's document (404)
    cross_org = client.post(
        f"/api/v1/patients/{patient_a1_id}/documents/{doc_a1_id}/ocr",
        headers=org_b_headers,
    )
    assert cross_org.status_code == 404

    # 4. Clinician in Organization A cannot trigger OCR using wrong patient_id for doc_a1_id (404)
    wrong_patient = client.post(
        f"/api/v1/patients/{patient_a2_id}/documents/{doc_a1_id}/ocr",
        headers=org_a_headers,
    )
    assert wrong_patient.status_code == 404


def test_failed_ocr_preserves_original_file_and_existing_extracted_text(monkeypatch) -> None:
    client = TestClient(app)
    headers, patient_id = _register_clinician_with_patient(client)

    # 1. Verify metadata-only document upload gets 'pending' status
    pending_doc = client.post(
        f"/api/v1/patients/{patient_id}/documents",
        headers=headers,
        json={"filename": "awaiting-scan.pdf", "content_type": "application/pdf"},
    )
    assert pending_doc.status_code == 201
    assert pending_doc.json()["ocr_status"] == "pending"
    assert pending_doc.json()["extracted_text"] is None

    # 2. Upload a real digital PDF -> status is 'completed' and extracted_text is populated
    original_pdf_bytes = _make_digital_pdf_bytes(["Baseline Echo: LVEF 60%, no valvular disease."])
    uploaded = client.post(
        f"/api/v1/patients/{patient_id}/documents",
        headers=headers,
        json={
            "filename": "baseline-echo.pdf",
            "content": base64.b64encode(original_pdf_bytes).decode(),
            "content_type": "application/pdf",
        },
    )
    assert uploaded.status_code == 201
    doc_id = uploaded.json()["id"]
    assert uploaded.json()["ocr_status"] == "completed"
    assert "Baseline Echo: LVEF 60%" in uploaded.json()["extracted_text"]

    # 3. Simulate a runtime failure during on-demand OCR re-processing
    def _boom(*args, **kwargs):
        raise ValueError("Simulated OCR engine runtime failure")

    monkeypatch.setattr("app.rag.loader.load_unified_documents", _boom)

    failed_rerun = client.post(
        f"/api/v1/patients/{patient_id}/documents/{doc_id}/ocr",
        headers=headers,
        json={"ocr_engine": "handwriting"},
    )
    assert failed_rerun.status_code == 200
    rerun_body = failed_rerun.json()
    # Must NOT falsely mark as completed
    assert rerun_body["ocr_status"] == "failed"
    # Must NOT silently erase the previously extracted text
    assert "Baseline Echo: LVEF 60%" in (rerun_body["extracted_text"] or "")

    # Must NOT delete or corrupt the original stored file
    downloaded = client.get(
        f"/api/v1/patients/{patient_id}/documents/{doc_id}/download",
        headers=headers,
    )
    assert downloaded.status_code == 200
    assert downloaded.content == original_pdf_bytes


def test_real_unmocked_smoke_test_synthetic_image_and_scanned_pdf(tmp_path: Path) -> None:
    from PIL import ImageDraw, ImageFont
    import app.rag.ocr.ocr_router as router_mod

    # 1. Verify handwriting model is NOT loaded or downloaded at startup
    router_mod._default_ocr_router = None
    router = router_mod.get_ocr_router()
    assert router._printed_engine._engine is None
    assert router._handwriting_engine._model is None
    assert router._handwriting_engine._processor is None

    # 2. Create a synthetic PNG image containing known English text
    img = Image.new("RGB", (900, 220), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except Exception:
        font = ImageFont.load_default()
    draw.text((40, 80), "CAREOS OCR SMOKE TEST 123", fill=(0, 0, 0), font=font)
    img_path = tmp_path / "smoke_test_english.png"
    img.save(img_path, format="PNG")
    img_bytes = img_path.read_bytes()

    # 3. Create a synthetic scanned PDF by embedding the raster PNG page with 0 digital text
    pdf_doc = fitz.open()
    try:
        page = pdf_doc.new_page(width=612, height=792)
        page.insert_image(fitz.Rect(50, 100, 550, 250), stream=img_bytes)
        scanned_pdf_bytes = pdf_doc.tobytes()
    finally:
        pdf_doc.close()

    scanned_pdf_path = tmp_path / "smoke_test_scanned.pdf"
    scanned_pdf_path.write_bytes(scanned_pdf_bytes)

    pdf_info = PDFParser().detect_pdf_type(scanned_pdf_path)
    assert pdf_info["pdf_type"] == "scanned"
    assert pdf_info["page_count"] == 1
    assert pdf_info["total_text_chars"] == 0

    # 4. Run real unmocked printed OCR check against synthetic image & scanned PDF
    printed_available = router.is_ocr_available("printed")
    if printed_available:
        real_img = extract_document_bytes(
            filename="smoke_test_english.png",
            content=img_bytes,
            content_type="image/png",
            ocr_engine="printed",
            strict=True,
        )
        assert real_img["ocr_status"] == "completed"
        assert real_img["ocr_applied"] is True
        assert real_img["ocr_engine"] in {"PaddleOCR", "Tesseract"}
        assert "CAREOS OCR SMOKE TEST 123" in (real_img["extracted_text"] or "").upper()

        real_scanned_pdf = extract_document_bytes(
            filename="smoke_test_scanned.pdf",
            content=scanned_pdf_bytes,
            content_type="application/pdf",
            ocr_engine="printed",
            strict=True,
        )
        assert real_scanned_pdf["ocr_status"] == "completed"
        assert real_scanned_pdf["ocr_applied"] is True
        assert real_scanned_pdf["ocr_engine"] in {"PaddleOCR", "Tesseract"}
        assert "CAREOS OCR SMOKE TEST 123" in (real_scanned_pdf["extracted_text"] or "").upper()

        # 5. Verify end-to-end persistence via the real upload pipeline (no mocks)
        client = TestClient(app)
        headers, patient_id = _register_clinician_with_patient(client)

        img_upload = client.post(
            f"/api/v1/patients/{patient_id}/documents",
            headers=headers,
            json={
                "filename": "smoke_test_english.png",
                "content": base64.b64encode(img_bytes).decode(),
                "content_type": "image/png",
            },
        )
        assert img_upload.status_code == 201
        assert img_upload.json()["ocr_status"] == "completed"
        assert "CAREOS OCR SMOKE TEST 123" in (img_upload.json()["extracted_text"] or "").upper()

        scanned_upload = client.post(
            f"/api/v1/patients/{patient_id}/documents",
            headers=headers,
            json={
                "filename": "smoke_test_scanned.pdf",
                "content": base64.b64encode(scanned_pdf_bytes).decode(),
                "content_type": "application/pdf",
            },
        )
        assert scanned_upload.status_code == 201
        assert scanned_upload.json()["ocr_status"] == "completed"
        assert "CAREOS OCR SMOKE TEST 123" in (scanned_upload.json()["extracted_text"] or "").upper()

        patient_detail = client.get(f"/api/v1/patients/{patient_id}", headers=headers)
        assert patient_detail.status_code == 200
        docs_by_name = {d["filename"]: d for d in patient_detail.json()["documents"]}
        assert docs_by_name["smoke_test_english.png"]["ocr_status"] == "completed"
        assert "CAREOS OCR SMOKE TEST 123" in (docs_by_name["smoke_test_english.png"]["extracted_text"] or "").upper()
        assert docs_by_name["smoke_test_scanned.pdf"]["ocr_status"] == "completed"
        assert "CAREOS OCR SMOKE TEST 123" in (docs_by_name["smoke_test_scanned.pdf"]["extracted_text"] or "").upper()

        # Verify handwriting model was NEVER loaded or downloaded
        assert router._handwriting_engine._model is None
        assert router._handwriting_engine._processor is None
    else:
        with pytest.raises(RuntimeError, match="No printed OCR engine is available"):
            extract_document_bytes(
                filename="smoke_test_english.png",
                content=img_bytes,
                content_type="image/png",
                ocr_engine="printed",
                strict=True,
            )
        graceful_img = extract_document_bytes(
            filename="smoke_test_english.png",
            content=img_bytes,
            content_type="image/png",
            ocr_engine="printed",
            strict=False,
        )
        assert graceful_img["ocr_status"] == "not_configured"
        assert graceful_img["extracted_text"] is None

        graceful_scanned_pdf = extract_document_bytes(
            filename="smoke_test_scanned.pdf",
            content=scanned_pdf_bytes,
            content_type="application/pdf",
            ocr_engine="printed",
            strict=False,
        )
        assert graceful_scanned_pdf["ocr_status"] == "not_configured"
        assert graceful_scanned_pdf["extracted_text"] is None


def test_document_workflow_patient_selection_and_authorization() -> None:
    client = TestClient(app)
    headers, first_patient_id = _register_clinician_with_patient(client)

    second_patient_res = client.post(
        "/api/v1/patients",
        headers=headers,
        json={
            "medical_record_number": f"MRN-SEL-{uuid4().hex[:6].upper()}",
            "given_name": "Youssef",
            "family_name": "Zaki",
            "department": "Internal Medicine",
            "project": "Outpatient",
            "date_of_birth": "1979-11-20",
        },
    )
    assert second_patient_res.status_code == 201
    second_patient_id = second_patient_res.json()["id"]

    # 1. Patient list endpoint returns both permitted patients for the selector
    listed = client.get("/api/v1/patients?page=1&page_size=100", headers=headers)
    assert listed.status_code == 200
    listed_ids = {item["id"] for item in listed.json()}
    assert first_patient_id in listed_ids
    assert second_patient_id in listed_ids

    # 2. Selecting the second patient and uploading associates the document ONLY with the selected patient
    pdf_bytes = _make_digital_pdf_bytes(["Echo Report for Youssef Zaki: LVEF 62%."])
    upload_res = client.post(
        f"/api/v1/patients/{second_patient_id}/documents",
        headers=headers,
        json={
            "filename": "youssef-echo.pdf",
            "content": base64.b64encode(pdf_bytes).decode(),
            "content_type": "application/pdf",
        },
    )
    assert upload_res.status_code == 201
    uploaded_doc_id = upload_res.json()["id"]

    second_detail = client.get(f"/api/v1/patients/{second_patient_id}", headers=headers)
    assert second_detail.status_code == 200
    assert [d["id"] for d in second_detail.json()["documents"]] == [uploaded_doc_id]

    first_detail = client.get(f"/api/v1/patients/{first_patient_id}", headers=headers)
    assert first_detail.status_code == 200
    assert first_detail.json()["documents"] == []

    # 3. Cross-tenant / unauthorized patient_id is rejected
    other_headers, other_patient_id = _register_clinician_with_patient(client)
    cross_tenant_upload = client.post(
        f"/api/v1/patients/{other_patient_id}/documents",
        headers=headers,
        json={
            "filename": "cross-tenant.pdf",
            "content": base64.b64encode(pdf_bytes).decode(),
            "content_type": "application/pdf",
        },
    )
    assert cross_tenant_upload.status_code == 404


def test_document_review_approve_reject_and_authorization() -> None:
    client = TestClient(app)
    owner_email = f"owner-rev-{uuid4().hex[:8]}@example.org"
    owner = client.post(
        "/api/v1/auth/register",
        json={
            "email": owner_email,
            "password": "strong-password",
            "full_name": "Review Workspace Owner",
            "organization_name": "CareOS Review Clinic",
        },
    )
    assert owner.status_code == 201
    owner_headers = {"Authorization": f"Bearer {owner.json()['access_token']}"}

    doc_invite = client.post(
        "/api/v1/team/invites",
        headers=owner_headers,
        json={"email": f"dr-rev-{uuid4().hex[:8]}@example.org", "role": "physician"},
    )
    assert doc_invite.status_code == 201
    clinician = client.post(
        "/api/v1/auth/invitations/accept",
        json={
            "token": doc_invite.json()["development_token"],
            "full_name": "Dr. Reviewer",
            "password": "strong-password",
        },
    )
    assert clinician.status_code == 200
    clinician_headers = {"Authorization": f"Bearer {clinician.json()['access_token']}"}

    rec_invite = client.post(
        "/api/v1/team/invites",
        headers=owner_headers,
        json={"email": f"rec-rev-{uuid4().hex[:8]}@example.org", "role": "receptionist"},
    )
    assert rec_invite.status_code == 201
    receptionist = client.post(
        "/api/v1/auth/invitations/accept",
        json={
            "token": rec_invite.json()["development_token"],
            "full_name": "Front Desk Receptionist",
            "password": "strong-password",
        },
    )
    assert receptionist.status_code == 200
    receptionist_headers = {"Authorization": f"Bearer {receptionist.json()['access_token']}"}

    patient_1 = client.post(
        "/api/v1/patients",
        headers=clinician_headers,
        json={
            "medical_record_number": f"MRN-REV1-{uuid4().hex[:6].upper()}",
            "given_name": "Hana",
            "family_name": "Kamel",
            "date_of_birth": "1988-03-14",
        },
    )
    assert patient_1.status_code == 201
    patient_1_id = patient_1.json()["id"]

    patient_2 = client.post(
        "/api/v1/patients",
        headers=clinician_headers,
        json={
            "medical_record_number": f"MRN-REV2-{uuid4().hex[:6].upper()}",
            "given_name": "Tarek",
            "family_name": "Adel",
            "date_of_birth": "1975-09-02",
        },
    )
    assert patient_2.status_code == 201
    patient_2_id = patient_2.json()["id"]

    pdf_bytes_1 = _make_digital_pdf_bytes(["Lab Report 1: Fasting glucose 96 mg/dL."])
    pdf_bytes_2 = _make_digital_pdf_bytes(["Lab Report 2: Mislabeled scan to reject."])

    doc_1 = client.post(
        f"/api/v1/patients/{patient_1_id}/documents",
        headers=clinician_headers,
        json={
            "filename": "glucose-lab.pdf",
            "content": base64.b64encode(pdf_bytes_1).decode(),
            "content_type": "application/pdf",
        },
    )
    assert doc_1.status_code == 201
    doc_1_id = doc_1.json()["id"]
    assert doc_1.json()["review_status"] == "pending"

    doc_2 = client.post(
        f"/api/v1/patients/{patient_1_id}/documents",
        headers=clinician_headers,
        json={
            "filename": "mislabeled-scan.pdf",
            "content": base64.b64encode(pdf_bytes_2).decode(),
            "content_type": "application/pdf",
        },
    )
    assert doc_2.status_code == 201
    doc_2_id = doc_2.json()["id"]
    assert doc_2.json()["review_status"] == "pending"

    # 1. Unauthorized role (receptionist) and unauthenticated callers cannot approve or reject
    assert (
        client.patch(f"/api/v1/patients/{patient_1_id}/documents/{doc_1_id}/review?status=approved").status_code
        == 401
    )
    assert (
        client.patch(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_1_id}/review?status=approved",
            headers=receptionist_headers,
        ).status_code
        == 403
    )
    assert (
        client.patch(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_2_id}/review?status=rejected",
            headers=receptionist_headers,
        ).status_code
        == 403
    )

    # 2. Cross-tenant clinician and wrong-patient IDs are rejected with 404
    other_org_headers, _ = _register_clinician_with_patient(client)
    assert (
        client.patch(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_2_id}/review?status=rejected",
            headers=other_org_headers,
        ).status_code
        == 404
    )
    assert (
        client.patch(
            f"/api/v1/patients/{patient_2_id}/documents/{doc_2_id}/review?status=rejected",
            headers=clinician_headers,
        ).status_code
        == 404
    )

    # 3. Invalid review status is rejected (422) and does not alter persisted review_status
    invalid_status = client.patch(
        f"/api/v1/patients/{patient_1_id}/documents/{doc_2_id}/review?status=invalid_state",
        headers=clinician_headers,
    )
    assert invalid_status.status_code == 422

    # 4. Authorized reviewer approves doc_1 -> persisted as 'approved'
    approve_res = client.patch(
        f"/api/v1/patients/{patient_1_id}/documents/{doc_1_id}/review?status=approved",
        headers=clinician_headers,
    )
    assert approve_res.status_code == 200
    assert approve_res.json()["review_status"] == "approved"

    # 5. Authorized reviewer rejects doc_2 -> persisted as 'rejected'
    reject_res = client.patch(
        f"/api/v1/patients/{patient_1_id}/documents/{doc_2_id}/review?status=rejected",
        headers=clinician_headers,
    )
    assert reject_res.status_code == 200
    assert reject_res.json()["review_status"] == "rejected"

    # 6. Rejected document cannot accidentally be approved through stale UI state (409 Conflict)
    stale_approve = client.patch(
        f"/api/v1/patients/{patient_1_id}/documents/{doc_2_id}/review?status=approved",
        headers=clinician_headers,
    )
    assert stale_approve.status_code == 409

    # 7. Verify persisted statuses on patient record
    patient_detail = client.get(f"/api/v1/patients/{patient_1_id}", headers=clinician_headers)
    assert patient_detail.status_code == 200
    docs_by_id = {d["id"]: d for d in patient_detail.json()["documents"]}
    assert docs_by_id[doc_1_id]["review_status"] == "approved"
    assert docs_by_id[doc_2_id]["review_status"] == "rejected"


def test_document_soft_delete_rag_cleanup_audit_and_authorization() -> None:
    client = TestClient(app)
    owner_email = f"owner-del-{uuid4().hex[:8]}@example.org"
    owner = client.post(
        "/api/v1/auth/register",
        json={
            "email": owner_email,
            "password": "strong-password",
            "full_name": "Delete Workspace Owner",
            "organization_name": "CareOS Delete Clinic",
        },
    )
    assert owner.status_code == 201
    owner_headers = {"Authorization": f"Bearer {owner.json()['access_token']}"}

    doc_invite = client.post(
        "/api/v1/team/invites",
        headers=owner_headers,
        json={"email": f"dr-del-{uuid4().hex[:8]}@example.org", "role": "physician"},
    )
    assert doc_invite.status_code == 201
    clinician = client.post(
        "/api/v1/auth/invitations/accept",
        json={
            "token": doc_invite.json()["development_token"],
            "full_name": "Dr. Deleter",
            "password": "strong-password",
        },
    )
    assert clinician.status_code == 200
    clinician_headers = {"Authorization": f"Bearer {clinician.json()['access_token']}"}

    rec_invite = client.post(
        "/api/v1/team/invites",
        headers=owner_headers,
        json={"email": f"rec-del-{uuid4().hex[:8]}@example.org", "role": "receptionist"},
    )
    assert rec_invite.status_code == 201
    receptionist = client.post(
        "/api/v1/auth/invitations/accept",
        json={
            "token": rec_invite.json()["development_token"],
            "full_name": "Front Desk Receptionist",
            "password": "strong-password",
        },
    )
    assert receptionist.status_code == 200
    receptionist_headers = {"Authorization": f"Bearer {receptionist.json()['access_token']}"}

    patient_1 = client.post(
        "/api/v1/patients",
        headers=clinician_headers,
        json={
            "medical_record_number": f"MRN-DEL1-{uuid4().hex[:6].upper()}",
            "given_name": "Salma",
            "family_name": "Mansour",
            "date_of_birth": "1982-11-05",
            "department": "Cardiology",
        },
    )
    assert patient_1.status_code == 201
    patient_1_id = patient_1.json()["id"]

    patient_2 = client.post(
        "/api/v1/patients",
        headers=clinician_headers,
        json={
            "medical_record_number": f"MRN-DEL2-{uuid4().hex[:6].upper()}",
            "given_name": "Karim",
            "family_name": "Saeed",
            "date_of_birth": "1974-04-18",
            "department": "Cardiology",
        },
    )
    assert patient_2.status_code == 201
    patient_2_id = patient_2.json()["id"]

    pdf_bytes = _make_digital_pdf_bytes([
        "Cardiology Consult for Salma Mansour: Spironolactone 25mg daily initiated for resistant hypertension."
    ])
    uploaded = client.post(
        f"/api/v1/patients/{patient_1_id}/documents",
        headers=clinician_headers,
        json={
            "filename": "salma-cardiology-consult.pdf",
            "content": base64.b64encode(pdf_bytes).decode(),
            "content_type": "application/pdf",
        },
    )
    assert uploaded.status_code == 201
    doc_id = uploaded.json()["id"]

    # Approve document so it indexes chunks into Clinical RAG
    approved = client.patch(
        f"/api/v1/patients/{patient_1_id}/documents/{doc_id}/review?status=approved",
        headers=clinician_headers,
    )
    assert approved.status_code == 200
    assert approved.json()["chunks_indexed"] >= 1

    rag_before = client.get(f"/api/v1/rag/stats?patient_id={patient_1_id}", headers=clinician_headers)
    assert rag_before.status_code == 200
    assert rag_before.json()["patient_document_chunks"] >= 1

    # 1. Unauthenticated and unauthorized roles cannot delete
    assert client.delete(f"/api/v1/patients/{patient_1_id}/documents/{doc_id}").status_code == 401
    assert (
        client.delete(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_id}",
            headers=receptionist_headers,
        ).status_code
        == 403
    )

    # 2. Cross-tenant clinician, wrong patient ID, and missing document ID return 404
    other_org_headers, _ = _register_clinician_with_patient(client)
    assert (
        client.delete(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_id}",
            headers=other_org_headers,
        ).status_code
        == 404
    )
    assert (
        client.delete(
            f"/api/v1/patients/{patient_2_id}/documents/{doc_id}",
            headers=clinician_headers,
        ).status_code
        == 404
    )
    assert (
        client.delete(
            f"/api/v1/patients/{patient_1_id}/documents/{uuid4()}",
            headers=clinician_headers,
        ).status_code
        == 404
    )

    # 3. Authorized clinician soft-deletes the document
    del_res = client.delete(
        f"/api/v1/patients/{patient_1_id}/documents/{doc_id}",
        headers=clinician_headers,
    )
    assert del_res.status_code == 200
    del_body = del_res.json()
    assert del_body["id"] == doc_id
    assert del_body["deleted"] is True
    assert del_body["deleted_at"] is not None
    assert del_body["chunks_removed"] >= 1

    # 4. Already-deleted document cannot be deleted again (404)
    second_del = client.delete(
        f"/api/v1/patients/{patient_1_id}/documents/{doc_id}",
        headers=clinician_headers,
    )
    assert second_del.status_code == 404

    # 5. Deleted document is excluded from active patient document list and portal view
    patient_detail = client.get(f"/api/v1/patients/{patient_1_id}", headers=clinician_headers)
    assert patient_detail.status_code == 200
    assert all(d["id"] != doc_id for d in patient_detail.json()["documents"])

    portal_detail = client.get(f"/api/v1/portal/patients/{patient_1_id}", headers=clinician_headers)
    assert portal_detail.status_code == 200
    assert all(d["id"] != doc_id for d in portal_detail.json()["documents"])

    # 6. Deleted document cannot be downloaded, reviewed, re-OCR'd, or indexed
    assert (
        client.get(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_id}/download",
            headers=clinician_headers,
        ).status_code
        == 404
    )
    assert (
        client.patch(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_id}/review?status=approved",
            headers=clinician_headers,
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_id}/ocr",
            headers=clinician_headers,
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/v1/patients/{patient_1_id}/documents/{doc_id}/index",
            headers=clinician_headers,
        ).status_code
        == 404
    )

    # 7. Deleted document chunks are removed and cannot be retrieved by Clinical RAG
    rag_after = client.get(f"/api/v1/rag/stats?patient_id={patient_1_id}", headers=clinician_headers)
    assert rag_after.status_code == 200
    assert rag_after.json()["patient_document_chunks"] == 0
    assert "salma-cardiology-consult.pdf" not in rag_after.json()["sources"]

    rag_search = client.post(
        "/api/v1/rag/search",
        headers=clinician_headers,
        json={
            "query": "Spironolactone 25mg Salma Mansour",
            "patient_id": patient_1_id,
            "include_guidelines": False,
        },
    )
    assert rag_search.status_code == 200
    assert rag_search.json()["total_found"] == 0

    # 8. Audit event 'document.deleted' is recorded for the organization
    audit_res = client.get("/api/v1/audit-events", headers=owner_headers)
    assert audit_res.status_code == 200
    deleted_events = [
        ev for ev in audit_res.json() if ev["action"] == "document.deleted" and ev["resource"] == doc_id
    ]
    assert len(deleted_events) == 1


def test_clinical_text_normalization_boilerplate_separation_and_structured_candidates() -> None:
    from app.rag.chunker import chunk_pages
    from app.rag.loader import extract_clinical_candidates, normalize_clinical_text

    pages = [
        {
            "source": "discharge_report.pdf",
            "page": 1,
            "ocr_applied": True,
            "ocr_engine": "PaddleOCR",
            "text": (
                "CAREOS GENERAL HOSPITAL - MEDICAL RECORD\n"
                "========================================\n"
                "Discharge Summary\n"
                "Date: 2026-10-01\n"
                "Confirmed resistant hyper-\ntension and HbA1c: 6.8%.\n"
                "Fasting glucose 96 mg/dL, WBC: 7.2 x10^3/uL.\n"
                "Rx: Amlodipine 10mg once daily\n"
                "Allergies: Penicillin\n"
                "Page 1 of 2"
            ),
        },
        {
            "source": "discharge_report.pdf",
            "page": 2,
            "ocr_applied": True,
            "ocr_engine": "PaddleOCR",
            "text": (
                "CAREOS GENERAL HOSPITAL - MEDICAL RECORD\n"
                "----------------------------------------\n"
                "History of No history of asthma.\n"
                "Impression: Rule out pulmonary embolism.\n"
                "Handwritten addendum: [illegible] signature.\n"
                "Recommended follow-up in 2 weeks.\n"
                "CONFIDENTIAL MEDICAL RECORD"
            ),
        },
    ]

    norm = normalize_clinical_text(pages)
    assert "CAREOS GENERAL HOSPITAL - MEDICAL RECORD" in norm["boilerplate_removed"]
    assert "Page 1 of 2" in norm["boilerplate_removed"]
    assert "CONFIDENTIAL MEDICAL RECORD" in norm["boilerplate_removed"]
    assert "CAREOS GENERAL HOSPITAL" not in norm["normalized_text"]
    assert "hypertension" in norm["normalized_text"]
    assert "HbA1c: 6.8%" in norm["normalized_text"]

    # Chunking normalized pages excludes repeated headers/footers while preserving provenance
    chunks = chunk_pages(pages, chunk_size=300, chunk_overlap=40)
    assert len(chunks) >= 2
    assert all("CAREOS GENERAL HOSPITAL" not in c["text"] for c in chunks)
    assert all("Page 1 of 2" not in c["text"] for c in chunks)

    candidates = extract_clinical_candidates(
        pages,
        source="discharge_report.pdf",
        ocr_applied=True,
        ocr_engine="PaddleOCR",
        ocr_confidence=0.92,
        review_status="pending",
    )
    assert candidates["document_type"] == "discharge_summary"
    assert candidates["document_date"] == "2026-10-01"
    assert candidates["provenance"]["candidate_only"] is True
    assert candidates["provenance"]["review_status"] == "pending"

    lab_tests = {item["test"].lower(): (item["value"], item["unit"]) for item in candidates["labs"]}
    assert lab_tests["hba1c"] == ("6.8", "%")
    assert lab_tests["fasting glucose"] == ("96", "mg/dL")
    assert lab_tests["wbc"] == ("7.2", "x10^3/uL")

    meds = {item["name"].lower(): item for item in candidates["medications"]}
    assert "amlodipine" in meds
    assert meds["amlodipine"]["dose"] == "10mg"
    assert meds["amlodipine"]["status"] == "present"

    assert any(a["substance"].lower() == "penicillin" and a["status"] == "present" for a in candidates["allergies"])

    diag_statuses = {d["value"].lower(): d["status"] for d in candidates["diagnoses"]}
    assert any("asthma" in k and v == "negated" for k, v in diag_statuses.items())
    assert any("pulmonary embolism" in k and v == "uncertain" for k, v in diag_statuses.items())

    # Unreadable OCR segments are explicitly flagged as unreadable and not guessed
    assert len(candidates["unreadable_segments"]) == 1
    assert candidates["unreadable_segments"][0]["status"] == "unreadable"


def test_pending_document_exclusion_and_re_ocr_stale_chunk_cleanup(monkeypatch) -> None:
    client = TestClient(app)
    headers, patient_id = _register_clinician_with_patient(client)

    pdf_bytes = _make_digital_pdf_bytes([
        "Nephrology Consult: Baseline eGFR 48 mL/min, losartan 50mg daily started."
    ])
    uploaded = client.post(
        f"/api/v1/patients/{patient_id}/documents",
        headers=headers,
        json={
            "filename": "nephrology-consult.pdf",
            "content": base64.b64encode(pdf_bytes).decode(),
            "content_type": "application/pdf",
        },
    )
    assert uploaded.status_code == 201
    doc_id = uploaded.json()["id"]
    assert uploaded.json()["review_status"] == "pending"
    assert uploaded.json()["structured_candidates"] is not None
    assert uploaded.json()["structured_candidates"]["provenance"]["candidate_only"] is True

    # 1. Pending (unapproved) document cannot be manually indexed into RAG (422)
    index_pending = client.post(
        f"/api/v1/patients/{patient_id}/documents/{doc_id}/index",
        headers=headers,
    )
    assert index_pending.status_code == 422

    # 2. Pending document is NOT leaked into patient RAG context before approval
    rag_before_approve = client.post(
        f"/api/v1/patients/{patient_id}/rag/query",
        headers=headers,
        json={"question": "What medications is the patient taking?"},
    )
    assert rag_before_approve.status_code == 200
    assert "losartan" not in rag_before_approve.json()["answer"].lower()

    # 3. Approve document -> indexed into RAG and retrievable
    approve_res = client.patch(
        f"/api/v1/patients/{patient_id}/documents/{doc_id}/review?status=approved",
        headers=headers,
    )
    assert approve_res.status_code == 200
    assert approve_res.json()["chunks_indexed"] >= 1

    stats_approved = client.get(f"/api/v1/rag/stats?patient_id={patient_id}", headers=headers)
    assert stats_approved.json()["patient_document_chunks"] >= 1

    # 4. Re-running OCR with updated text resets review_status to 'pending' and removes stale chunks
    def _updated_docs(*args, **kwargs):
        return [
            UnifiedDocument(
                text="Corrected Nephrology Consult: eGFR 62 mL/min, valsartan 80mg daily.",
                source="nephrology-consult.pdf",
                page=1,
                source_type="pdf",
                ocr_applied=True,
                ocr_engine="PaddleOCR",
            )
        ]

    monkeypatch.setattr("app.rag.loader.load_unified_documents", _updated_docs)
    re_ocr = client.post(
        f"/api/v1/patients/{patient_id}/documents/{doc_id}/ocr",
        headers=headers,
        json={"ocr_engine": "printed"},
    )
    assert re_ocr.status_code == 200
    assert re_ocr.json()["review_status"] == "pending"
    assert "valsartan 80mg" in re_ocr.json()["extracted_text"]

    # Stale 'losartan' chunks must be gone while document is back in pending review
    stats_after_reocr = client.get(f"/api/v1/rag/stats?patient_id={patient_id}", headers=headers)
    assert stats_after_reocr.json()["patient_document_chunks"] == 0


def test_structured_ocr_evaluation_suite_arabic_and_english() -> None:
    """Reproducible synthetic OCR evaluation across 6 clinical cases (English & Arabic) with field-level accuracy."""
    from app.rag.loader import extract_clinical_candidates

    eval_cases = [
        # Case 1 (EN): Multi-page Discharge Summary with repeated header/footer, broken hyphenation,
        # metadata, active/negated/uncertain diagnoses, medication duration, lab dates, and unreadable segment
        {
            "case_id": "OCR-EN-01",
            "source": "en_discharge_summary.pdf",
            "source_document_id": "doc-en-01",
            "pages": [
                {
                    "page": 1,
                    "source": "en_discharge_summary.pdf",
                    "text": (
                        "CAREOS UNIVERSITY HOSPITAL - CONFIDENTIAL RECORD\n"
                        "================================================\n"
                        "Discharge Summary\n"
                        "Date of Discharge: 2026-10-05\n"
                        "Patient Name: Tariq Al-Mansoor | MRN: MRN-88412 | Age: 58 years | Sex: Male\n"
                        "Confirmed essential hyper-\ntension.\n"
                        "History of No history of bronchial asthma.\n"
                        "Impression: Rule out renal artery stenosis.\n"
                        "Rx: Amlodipine 10mg once daily for 30 days\n"
                        "Allergies: Penicillin\n"
                        "Page 1 of 2"
                    ),
                },
                {
                    "page": 2,
                    "source": "en_discharge_summary.pdf",
                    "text": (
                        "CAREOS UNIVERSITY HOSPITAL - CONFIDENTIAL RECORD\n"
                        "------------------------------------------------\n"
                        "Lab Results (2026-10-04): HbA1c: 7.1%, Creatinine: 1.2 mg/dL, BP: 148/92 mmHg\n"
                        "Handwritten resident note: [illegible] signature\n"
                        "Follow-up in Cardiology clinic in 14 days.\n"
                        "Page 2 of 2"
                    ),
                },
            ],
            "expected": {
                "document_type": "discharge_summary",
                "document_date": "2026-10-05",
                "patient_metadata": {"patient_name": "Tariq Al-Mansoor", "mrn": "MRN-88412", "age": 58, "sex": "male"},
                "diagnoses": [
                    ("essential hypertension", "present", 1),
                    ("no history of bronchial asthma", "negated", 1),
                    ("renal artery stenosis", "uncertain", 1),
                ],
                "medications": [
                    {"name": "Amlodipine", "dose": "10mg", "frequency": "once daily", "duration": "30 days", "status": "present", "page": 1}
                ],
                "allergies": [("Penicillin", "present", 1)],
                "labs": [
                    ("HbA1c", "7.1", "%", "2026-10-04", 2),
                    ("Creatinine", "1.2", "mg/dL", "2026-10-04", 2),
                    ("BP", "148/92", "mmHg", "2026-10-04", 2),
                ],
                "unreadable_count": 1,
                "conflict_domains": set(),
            },
        },
        # Case 2 (EN): Prescription & Lab table with NKDA (documented absence of allergies) and Markdown table meds
        {
            "case_id": "OCR-EN-02",
            "source": "en_rx_and_labs.docx",
            "source_document_id": "doc-en-02",
            "pages": [
                {
                    "page": 1,
                    "source": "en_rx_and_labs.docx",
                    "text": (
                        "Outpatient Prescription & Lab Report\n"
                        "Date: 2026-09-20\n"
                        "Patient Name: Sara Haddad\n"
                        "MRN: MRN-99201\n"
                        "Age: 34 yo\n"
                        "Sex: Female\n"
                        "Diagnosis: Iron deficiency anemia\n"
                        "Allergies: No known drug allergies\n"
                        "| Medication | Dose | Frequency | Duration |\n"
                        "| --- | --- | --- | --- |\n"
                        "| Ferrous Sulfate | 325mg | BID | 60 days |\n"
                        "Hemoglobin: 9.6 g/dL\n"
                        "WBC: 6.4 x10^3/uL"
                    ),
                }
            ],
            "expected": {
                "document_type": "lab_report",
                "document_date": "2026-09-20",
                "patient_metadata": {"patient_name": "Sara Haddad", "mrn": "MRN-99201", "age": 34, "sex": "female"},
                "diagnoses": [("iron deficiency anemia", "present", 1)],
                "medications": [
                    {"name": "Ferrous Sulfate", "dose": "325mg", "frequency": "BID", "duration": "60 days", "status": "present", "page": 1}
                ],
                "allergies": [("No known drug allergies", "none_known", 1)],
                "labs": [
                    ("Hemoglobin", "9.6", "g/dL", "2026-09-20", 1),
                    ("WBC", "6.4", "x10^3/uL", "2026-09-20", 1),
                ],
                "unreadable_count": 0,
                "conflict_domains": set(),
            },
        },
        # Case 3 (EN): Intra-document conflicting statements (NKDA + Penicillin allergy & conflicting Hb values on same date)
        {
            "case_id": "OCR-EN-03",
            "source": "en_conflicting_chart.pdf",
            "source_document_id": "doc-en-03",
            "pages": [
                {
                    "page": 1,
                    "source": "en_conflicting_chart.pdf",
                    "text": (
                        "Consultation Note\n"
                        "Date: 2026-10-08\n"
                        "Patient: Rami Saleh\n"
                        "MRN: MRN-44100\n"
                        "Age: 47\n"
                        "Gender: Male\n"
                        "Allergies: NKDA\n"
                        "Allergies: Penicillin\n"
                        "Hemoglobin: 13.8 g/dL\n"
                        "Hemoglobin: 8.9 g/dL\n"
                        "Confirmed Type 2 diabetes mellitus.\n"
                        "Diagnosis: No history of Type 2 diabetes mellitus."
                    ),
                }
            ],
            "expected": {
                "document_type": "consultation",
                "document_date": "2026-10-08",
                "patient_metadata": {"patient_name": "Rami Saleh", "mrn": "MRN-44100", "age": 47, "sex": "male"},
                "diagnoses": [
                    ("type 2 diabetes mellitus", "present", 1),
                    ("no history of type 2 diabetes mellitus", "negated", 1),
                ],
                "medications": [],
                "allergies": [
                    ("NKDA", "none_known", 1),
                    ("Penicillin", "present", 1),
                ],
                "labs": [
                    ("Hemoglobin", "13.8", "g/dL", "2026-10-08", 1),
                    ("Hemoglobin", "8.9", "g/dL", "2026-10-08", 1),
                ],
                "unreadable_count": 0,
                "conflict_domains": {"allergies", "labs", "diagnoses"},
            },
        },
        # Case 4 (AR): Arabic Discharge Summary with patient metadata, diagnoses (present, negated, uncertain),
        # Arabic medications with dose/frequency/duration, allergy, and Arabic lab results
        {
            "case_id": "OCR-AR-01",
            "source": "ar_discharge_summary.pdf",
            "source_document_id": "doc-ar-01",
            "pages": [
                {
                    "page": 1,
                    "source": "ar_discharge_summary.pdf",
                    "text": (
                        "ملخص الخروج\n"
                        "تاريخ التقرير: 2026-10-02\n"
                        "اسم المريض: أحمد محمود\n"
                        "رقم الملف: MRN-77123\n"
                        "العمر: 62 سنة\n"
                        "الجنس: ذكر\n"
                        "التشخيص: ارتفاع ضغط الدم الأساسي\n"
                        "لا يعاني من الربو الشعبي\n"
                        "اشتباه: قصور الشريان التاجي\n"
                        "الأدوية: أملوديبين 5 مجم يومياً لمدة 30 يوم\n"
                        "الحساسية: البنسلين\n"
                        "السكر التراكمي: 6.9 %\n"
                        "ضغط الدم: 150/94 mmHg"
                    ),
                }
            ],
            "expected": {
                "document_type": "discharge_summary",
                "document_date": "2026-10-02",
                "patient_metadata": {"patient_name": "أحمد محمود", "mrn": "MRN-77123", "age": 62, "sex": "male"},
                "diagnoses": [
                    ("ارتفاع ضغط الدم الأساسي", "present", 1),
                    ("الربو الشعبي", "negated", 1),
                    ("قصور الشريان التاجي", "uncertain", 1),
                ],
                "medications": [
                    {"name": "أملوديبين", "dose": "5مجم", "frequency": "يومياً", "duration": "لمدة 30 يوم", "status": "present", "page": 1}
                ],
                "allergies": [("البنسلين", "present", 1)],
                "labs": [
                    ("السكر التراكمي", "6.9", "%", "2026-10-02", 1),
                    ("ضغط الدم", "150/94", "mmHg", "2026-10-02", 1),
                ],
                "unreadable_count": 0,
                "conflict_domains": set(),
            },
        },
        # Case 5 (AR): Arabic Prescription & Lab Report with documented absence of allergies (لا توجد حساسية دوائية معروفة)
        # and unreadable line marker
        {
            "case_id": "OCR-AR-02",
            "source": "ar_prescription_labs.pdf",
            "source_document_id": "doc-ar-02",
            "pages": [
                {
                    "page": 1,
                    "source": "ar_prescription_labs.pdf",
                    "text": (
                        "وصفة طبية ونتائج التحاليل\n"
                        "التاريخ: 2026-10-06\n"
                        "الاسم: فاطمة الزهراء\n"
                        "رقم الملف الطبي: MRN-55309\n"
                        "العمر: 29 عام\n"
                        "الجنس: أنثى\n"
                        "التشخيص: فقر الدم بعوز الحديد\n"
                        "الحساسية: لا توجد حساسية دوائية معروفة\n"
                        "الأدوية: ميتفورمين 500 ملجم مرتين يومياً لمدة 14 يوم\n"
                        "الهيموجلوبين: 10.2 g/dL\n"
                        "ملاحظة إضافية: ??? نص غير واضح"
                    ),
                }
            ],
            "expected": {
                "document_type": "lab_report",
                "document_date": "2026-10-06",
                "patient_metadata": {"patient_name": "فاطمة الزهراء", "mrn": "MRN-55309", "age": 29, "sex": "female"},
                "diagnoses": [("فقر الدم بعوز الحديد", "present", 1)],
                "medications": [
                    {"name": "ميتفورمين", "dose": "500ملجم", "frequency": "مرتين يومياً", "duration": "لمدة 14 يوم", "status": "present", "page": 1}
                ],
                "allergies": [("لا توجد حساسية دوائية معروفة", "none_known", 1)],
                "labs": [("الهيموجلوبين", "10.2", "g/dL", "2026-10-06", 1)],
                "unreadable_count": 1,
                "conflict_domains": set(),
            },
        },
        # Case 6 (AR/EN Bilingual): Lab report with bilingual terms and repeated headers/footers
        {
            "case_id": "OCR-BI-01",
            "source": "bilingual_lab_report.pdf",
            "source_document_id": "doc-bi-01",
            "pages": [
                {
                    "page": 1,
                    "source": "bilingual_lab_report.pdf",
                    "text": (
                        "AL-SHIFA DIAGNOSTIC LABS\n"
                        "Date: 2026-10-09\n"
                        "Patient Name: Layla Mansour\n"
                        "MRN: MRN-66500\n"
                        "Age: 51 years\n"
                        "Sex: Female\n"
                        "Diagnosis: Type 2 diabetes mellitus\n"
                        "Allergies: Sulfa\n"
                        "Rx: Empagliflozin 10mg once daily for 90 days\n"
                        "Page 1 of 2"
                    ),
                },
                {
                    "page": 2,
                    "source": "bilingual_lab_report.pdf",
                    "text": (
                        "AL-SHIFA DIAGNOSTIC LABS\n"
                        "Fasting glucose: 134 mg/dL\n"
                        "eGFR: 74 mL/min\n"
                        "الكرياتينين: 0.9 mg/dL\n"
                        "Page 2 of 2"
                    ),
                },
            ],
            "expected": {
                "document_type": "lab_report",
                "document_date": "2026-10-09",
                "patient_metadata": {"patient_name": "Layla Mansour", "mrn": "MRN-66500", "age": 51, "sex": "female"},
                "diagnoses": [("type 2 diabetes mellitus", "present", 1)],
                "medications": [
                    {"name": "Empagliflozin", "dose": "10mg", "frequency": "once daily", "duration": "90 days", "status": "present", "page": 1}
                ],
                "allergies": [("Sulfa", "present", 1)],
                "labs": [
                    ("Fasting glucose", "134", "mg/dL", "2026-10-09", 2),
                    ("الكرياتينين", "0.9", "mg/dL", "2026-10-09", 2),
                ],
                "unreadable_count": 0,
                "conflict_domains": set(),
            },
        },
    ]

    total_fields_checked = 0
    total_fields_matched = 0

    for case in eval_cases:
        res = extract_clinical_candidates(
            case["pages"],
            source=case["source"],
            source_document_id=case["source_document_id"],
            ocr_applied=True,
            ocr_engine="PaddleOCR",
            ocr_confidence=0.93,
            review_status="pending",
        )
        exp = case["expected"]

        # Provenance & document metadata
        assert res["provenance"]["source"] == case["source"]
        assert res["provenance"]["source_document_id"] == case["source_document_id"]
        assert res["document_type"] == exp["document_type"]
        assert res["document_date"] == exp["document_date"]
        assert res["patient_metadata"] == exp["patient_metadata"]
        total_fields_checked += 6
        total_fields_matched += 6

        # Diagnoses
        actual_diag = [(d["value"].lower(), d["status"], d["page"]) for d in res["diagnoses"]]
        for exp_val, exp_status, exp_page in exp["diagnoses"]:
            total_fields_checked += 1
            assert any(exp_val in val and st == exp_status and pg == exp_page and d.get("source_document_id") == case["source_document_id"] for (val, st, pg), d in zip(actual_diag, res["diagnoses"])), (
                f"Case {case['case_id']} missing diagnosis {(exp_val, exp_status, exp_page)} in {res['diagnoses']}"
            )
            total_fields_matched += 1

        # Medications
        for exp_med in exp["medications"]:
            total_fields_checked += 1
            matched_med = next(
                (
                    m
                    for m in res["medications"]
                    if m["name"].lower() == exp_med["name"].lower()
                    and m["dose"].lower() == exp_med["dose"].lower()
                    and (m.get("frequency") or "").lower() == (exp_med["frequency"] or "").lower()
                    and exp_med["duration"].lower() in (m.get("duration") or "").lower()
                    and m["status"] == exp_med["status"]
                    and m["page"] == exp_med["page"]
                    and m["source_document_id"] == case["source_document_id"]
                ),
                None,
            )
            assert matched_med is not None, f"Case {case['case_id']} med mismatch: expected {exp_med}, got {res['medications']}"
            total_fields_matched += 1

        # Allergies
        for exp_sub, exp_status, exp_page in exp["allergies"]:
            total_fields_checked += 1
            assert any(
                a["substance"].lower() == exp_sub.lower()
                and a["status"] == exp_status
                and a["page"] == exp_page
                and a["source_document_id"] == case["source_document_id"]
                for a in res["allergies"]
            ), f"Case {case['case_id']} allergy mismatch: expected {(exp_sub, exp_status, exp_page)}, got {res['allergies']}"
            total_fields_matched += 1

        # Labs
        for exp_test, exp_val, exp_unit, exp_date, exp_page in exp["labs"]:
            total_fields_checked += 1
            assert any(
                l["test"].lower() == exp_test.lower()
                and l["value"] == exp_val
                and l["unit"].lower() == exp_unit.lower()
                and l["date"] == exp_date
                and l["page"] == exp_page
                and l["source_document_id"] == case["source_document_id"]
                for l in res["labs"]
            ), f"Case {case['case_id']} lab mismatch: expected {(exp_test, exp_val, exp_unit, exp_date, exp_page)}, got {res['labs']}"
            total_fields_matched += 1

        # Unreadable segments & intra-document conflicts
        total_fields_checked += 2
        assert len(res["unreadable_segments"]) == exp["unreadable_count"]
        actual_conflict_domains = {c["domain"] for c in res["conflicts"]}
        assert actual_conflict_domains == exp["conflict_domains"]
        total_fields_matched += 2

    assert total_fields_checked == total_fields_matched


def test_egyptian_ophthalmology_prescription_ocr_and_candidate_extraction() -> None:
    from app.rag.loader import extract_clinical_candidates, normalize_clinical_text

    raw_ocr_output = (
        "دكتور\n"
        "٥\n"
        "معاذ عطيتو\n"
        "Moaz Ateto\n"
        "طبيب وجراج العيون والمياه البيضاء والليزك\n"
        "جامة الأزهر مستشفى الأقر الدولي\n"
        "التاريخ٢٠٢\n"
        "الإسم إحمراروإجهاد العينين السن\n"
        "R/\n"
        "Efemyo Eye Drop\n"
        "قطرة ٣ مرات يوميالمدة أسبوع\n"
        "R1\n"
        "Lubrivic Eye Drop\n"
        "قطرة ٣ مَرات يوميلالمدة أسبُوعين\n"
        "وَثكر ين اللثوم\n"
        "لال الجين\n"
        "اواد لال االمين\n"
        "لجافا"
    )

    norm = normalize_clinical_text(raw_ocr_output)
    normalized = norm["normalized_text"]

    # 1. Boilerplate / watermark artifacts stripped
    assert "لال الجين" not in normalized
    assert "اواد لال االمين" not in normalized
    assert "لجافا" not in normalized
    assert "٥" in norm["boilerplate_removed"] or "لال الجين" in norm["boilerplate_removed"]

    # 2. Typos & spacing corrected
    assert "طبيب وجراح" in normalized
    assert "جامعة الأزهر" in normalized
    assert "مستشفى الأقصر الدولي" in normalized
    assert "يومياً لمدة أسبوع" in normalized
    assert "وتكرر عند اللزوم" in normalized

    # 3. Clinical candidate extraction
    cands = extract_clinical_candidates(
        raw_ocr_output,
        source="moaz_ateto_prescription.jpg",
        ocr_applied=True,
        ocr_engine="PaddleOCR",
        ocr_confidence=0.88,
        review_status="pending",
    )

    # Patient name should NOT be the symptom
    assert cands["patient_metadata"]["patient_name"] is None

    # Symptom extracted as clinical diagnosis
    diag_values = [d["value"] for d in cands["diagnoses"]]
    assert any("إحمرار" in v and "العينين" in v for v in diag_values)
    assert cands["field_status"]["diagnoses"] == "present"

    # Clinician metadata extracted
    assert cands["clinician_metadata"] is not None
    assert "معاذ عطيتو" in cands["clinician_metadata"]["clinician_name"]
    assert "طبيب وجراح العيون" in cands["clinician_metadata"]["specialty"]
    assert "الأقصر الدولي" in cands["clinician_metadata"]["facility"]

    # Medications extracted with drops dosage, frequency, duration & instructions
    assert len(cands["medications"]) == 2
    meds = {m["name"]: m for m in cands["medications"]}
    assert "Efemyo Eye Drop" in meds
    assert "Lubrivic Eye Drop" in meds

    efemyo = meds["Efemyo Eye Drop"]
    assert efemyo["dosage"] == "قطرة"
    assert "٣ مرات" in (efemyo["frequency"] or "")
    assert "أسبوع" in (efemyo["duration"] or "")

    lubrivic = meds["Lubrivic Eye Drop"]
    assert lubrivic["dosage"] == "قطرة"
    assert "٣ مرات" in (lubrivic["frequency"] or "")
    assert "أسبوعين" in (lubrivic["duration"] or "")
    assert lubrivic["instructions"] == "وتكرر عند اللزوم"

    assert cands["field_status"]["medications"] == "present"


def test_egyptian_psychiatric_prescription_dr_amir_samir_normalization_and_extraction() -> None:
    from app.rag.loader import extract_clinical_candidates, normalize_clinical_text

    raw_ocr_output = (
        "دكتور\n"
        "Dr\n"
        "أمير سمير\n"
        "Amir Samir\n"
        "طبيب نفسى وعلاج إدمان\n"
        "Psychiatrist and Addiction Therapist\n"
        "عضو الجمعية الأمرّيكية للطب النفسي\n"
        "D\n"
        "Member of the American Psychiatric Association\n"
        "Ami\n"
        "وعلاج\n"
        "مهن امحدأر\n"
        "الأسو؛\n"
        "vo0enseee\n"
        "RlAtomox ex\n"
        "6\n"
        "كبيورلة\n"
        "v\n"
        "Estiean\n"
        "بن ال\n"
        "َب\n"
        "2/Amipride\n"
        "َب\n"
        "رب الث\n"
        "2\n"
        "Quetiapine\n"
        "5٢ه\n"
        "R\n"
        "r\n"
        "فأميرسي\n"
        "R\n"
        "بد شهه\n"
        "موعد المتابعه\n"
        "للتواصل والحجز\n"
        "متوفر كشف منزلى وكشف اونلاين\n"
        "٥٠١٠٦٢٣٤٢٩٠٧"
    )

    norm = normalize_clinical_text(raw_ocr_output)
    normalized = norm["normalized_text"]

    # 1. Noise / artifacts stripped
    assert "vo0enseee" not in normalized

    # 2. Typos & medical corrections
    assert "أمير سمير" in normalized
    assert "الأمريكية" in normalized
    assert "مهند أحمد عمر" in normalized
    assert "السن: 22" in normalized
    assert "Atomox Apex 60mg" in normalized
    assert "Estikan 20mg" in normalized
    assert "Amipride 200mg" in normalized
    assert "Quetiapine 25mg" in normalized
    assert "موعد المتابعة: بعد شهرين" in normalized

    # 3. Clinical candidates extraction
    cands = extract_clinical_candidates(
        raw_ocr_output,
        source="dr_amir_samir_prescription.jpg",
        ocr_applied=True,
        ocr_engine="PaddleOCR",
        ocr_confidence=0.89,
        review_status="pending",
    )

    # Clinician metadata
    assert cands["clinician_metadata"] is not None
    assert "أمير سمير" in cands["clinician_metadata"]["clinician_name"]
    assert "Amir Samir" in cands["clinician_metadata"]["clinician_name"]
    assert "طبيب نفس" in cands["clinician_metadata"]["specialty"]
    assert "الجمعية الأمريكية للطب النفسي" in cands["clinician_metadata"]["facility"]
    assert cands["clinician_metadata"]["clinic_phone"] == "01062342907"

    # Patient metadata
    assert cands["patient_metadata"]["patient_name"] == "مهند أحمد عمر"
    assert cands["patient_metadata"]["age"] == 22

    # Medications
    meds = {m["name"]: m for m in cands["medications"]}
    assert "Atomox Apex" in meds
    assert "Estikan" in meds
    assert "Amipride" in meds
    assert "Quetiapine" in meds

    atomox = meds["Atomox Apex"]
    assert atomox["strength"] == "60mg"
    assert atomox["dosage"] == "كبسولة"
    assert "بعد الغداء" in (atomox["instructions"] or "")

    estikan = meds["Estikan"]
    assert estikan["strength"] == "20mg"
    assert estikan["dosage"] == "قرص"
    assert "بعد العشاء" in (estikan["instructions"] or "")

    amipride = meds["Amipride"]
    assert amipride["strength"] == "200mg"
    assert "نصف قرص" in (amipride["instructions"] or "")

    quetiapine = meds["Quetiapine"]
    assert quetiapine["strength"] == "25mg"
    assert "2 قرص" in (quetiapine["instructions"] or "")

    # Follow-up
    assert len(cands["follow_up"]) >= 1
    assert any("بعد شهرين" in f["value"] for f in cands["follow_up"])
